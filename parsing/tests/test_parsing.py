"""Standalone document-processing contract, without importing a RAG backend."""

import pytest
from docx import Document
from pdfgen import write_pdf

from parsing.chunker import chunk_document, split_text
from parsing.cleaner import normalize_text
from parsing.parser import parse_document

# PDF fragments of 80 characters or fewer are dropped as headers, so page texts here are longer.
PAGE_1 = "Short useful text: warm up for five minutes, then do three sets of ten slow squats."
PAGE_2 = "Second page: rest for two minutes between sets and keep the back straight at all times."


@pytest.mark.parametrize(("raw", "expected"), [
    ("test training text t n", "test training text t n"),
    ("test\t\ttext  training", "test text training"),
    ("first\r\nsecond\rthird", "first\nsecond\nthird"),
    ("first\n\n\n\nsecond", "first\n\nsecond"),
    ("first paragraph\n\nsecond paragraph", "first paragraph\n\nsecond paragraph"),
    ("\x00 first \n \t second \x00", "first\nsecond"),
    (r"literal\t and literal\n", r"literal\t and literal\n"),
])
def test_cleaning(raw, expected):
    assert normalize_text(raw) == expected
    assert normalize_text(expected) == expected


@pytest.mark.parametrize(("separator", "expected"), [
    ("\n\n", "first words"), ("\n", "first words"),
    (". ", "first words."), (" ", "first words second"),
])
def test_split_boundaries(separator, expected):
    assert split_text("first words" + separator + "second longword finalword", 23, 3)[0] == expected


def test_paragraph_priority_and_word_overlap():
    text = "first paragraph\n\nsecond line. more words to continue the document"
    assert split_text(text, 40, 5)[0] == "first paragraph"
    words = [f"word{i:02}" for i in range(30)]
    chunks = split_text(" ".join(words), 40, 12)
    assert {word for chunk in chunks for word in chunk.split()} == set(words)
    assert all(set(a.split()) & set(b.split()) for a, b in zip(chunks, chunks[1:]))


def test_pdf_pages_and_chunk_metadata(tmp_path):
    path = write_pdf(tmp_path / "guide.PDF", [PAGE_1, PAGE_2])
    parsed = parse_document(path)
    assert [p["page"] for p in parsed] == [1, 2]
    assert all(p["source"] == "guide.PDF" and p["file_type"] == "pdf" for p in parsed)
    chunks = chunk_document(path, corpus_root=tmp_path)
    assert [c["page"] for c in chunks] == [1, 2]
    assert [c["text"] for c in chunks] == [PAGE_1, PAGE_2]
    assert all(c["file_type"] == "pdf" and c["chunk_index"] == 0 for c in chunks)
    assert len({c["id"] for c in chunks}) == 2


def test_docx_tables_metadata_unique_stable_ids(tmp_path):
    document = Document()
    document.add_paragraph("test training " * 90)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "exercise"
    table.cell(0, 1).text = "table contents"
    document.add_paragraph("final paragraph")
    all_chunks = []
    for directory in (tmp_path, tmp_path / "nested"):
        directory.mkdir(exist_ok=True)
        path = directory / "notes.docx"
        document.save(path)
        [parsed] = parse_document(path)
        assert parsed["page"] is None
        assert parsed["source"] == "notes.docx" and parsed["file_type"] == "docx"
        assert "exercise\ttable contents" in parsed["text"]
        assert parsed["text"].index("test training") < parsed["text"].index("exercise") < parsed["text"].index("final paragraph")
        chunks = chunk_document(path, corpus_root=tmp_path)
        assert chunks == chunk_document(path, corpus_root=tmp_path)
        assert len(chunks) > 1
        assert [c["chunk_index"] for c in chunks] == list(range(len(chunks)))
        assert all(c["page"] is None and c["file_type"] == "docx" for c in chunks)
        assert all(c["source"] == path.relative_to(tmp_path).as_posix() for c in chunks)
        assert all(c["title"] == "notes.docx" for c in chunks)
        assert any("exercise table contents" in c["text"] for c in chunks)
        all_chunks.extend(chunks)
    assert len({c["id"] for c in all_chunks}) == len(all_chunks)


def test_who_compatibility(tmp_path):
    path = write_pdf(tmp_path / "who_2020_physical_activity.pdf", [PAGE_1])
    [chunk] = chunk_document(path, corpus_root=tmp_path)
    assert chunk["source"] == "who2020"
    assert chunk["id"] == "who:400:1:0"
    assert chunk["file_type"] == "pdf" and chunk["page"] == 1
