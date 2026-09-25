"""Regression tests from the parsing review: splitter invariants, extraction edge cases, index integration."""

import logging
import random

import pytest
from docx import Document
from pdfgen import write_pdf

from parsing.chunker import MIN_PDF_CHUNK_CHARS, PARSING_VERSION, _spans, chunk_document, split_text
from parsing.cleaner import normalize_text
from parsing.parser import parse_document

LONG = "Keep the back straight, brace the core and lower the hips until the thighs are parallel to the floor."


def _random_text(rnd: random.Random, max_word: int) -> str:
    parts = []
    for i in range(rnd.randint(0, 150)):
        parts.append(f"w{i:03d}" + "x" * rnd.randint(0, max_word - 4))
        parts.append(rnd.choice([" ", " ", " ", "\n", "\n\n", ". ", " \t "]))
    return "".join(parts)


@pytest.mark.parametrize("seed", range(40))
def test_split_spans_cover_text_without_gaps_or_duplicates(seed):
    rnd = random.Random(seed)
    size = rnd.choice([10, 23, 40, 100, 400])
    overlap = rnd.choice([0, 1, size // 4, size // 2, size - 2, size - 1])
    text = normalize_text(_random_text(rnd, max_word=rnd.choice([6, 12, 60, 900])))
    if len(text) <= size:
        return
    spans = _spans(text, size, overlap)
    assert spans[0][0] == 0 and spans[-1][1] == len(text)
    for (s1, e1), (s2, e2) in zip(spans, spans[1:]):
        assert s1 < s2 <= e1 < e2  # no gap, and each chunk reaches past the previous one
    assert all(end - start <= size for start, end in spans)
    assert len(spans) <= len(text)


@pytest.mark.parametrize("seed", range(40))
def test_split_keeps_every_word_whole_and_in_order(seed):
    rnd = random.Random(1000 + seed)
    size = rnd.choice([40, 100, 400])
    overlap = rnd.choice([0, size // 4, size // 2])
    text = _random_text(rnd, max_word=12)
    words = normalize_text(text).split()
    chunks = split_text(text, size, overlap)
    seen = list(dict.fromkeys(w for chunk in chunks for w in chunk.split()))
    assert seen == list(dict.fromkeys(words))


def test_split_does_not_repeat_the_overlap_as_its_own_chunk():
    text = ("Warm up for five minutes before lifting.\n\n"
            "Keep the back straight and brace the core during every repetition of the squat")
    chunks = split_text(text, 60, 20)
    assert not any(b in a for a, b in zip(chunks, chunks[1:]))
    assert chunks[0] == "Warm up for five minutes before lifting."


def test_split_edge_cases():
    assert split_text(" \n\t \xa0 \n\n ", 40, 10) == []
    word = "a" * 95
    pieces = split_text(word, 40, 10)
    assert all(len(p) <= 40 for p in pieces) and "".join(pieces[:1]) == word[:40]
    assert set("".join(pieces)) == {"a"}
    extreme = split_text(" ".join(f"word{i}" for i in range(40)), 30, 29)
    assert extreme and len(extreme) <= 40  # overlap close to size still terminates, one word per step at most
    with pytest.raises(ValueError):
        split_text("text", 10, 10)


@pytest.mark.parametrize(("raw", "expected"), [
    ("and\xa0no increase", "and no increase"),
    ("physi­cal ac​tivity﻿", "physical activity"),
    ("ﬁtness eﬀort", "fitness effort"),
    ("page one\fpage two", "page one\npage two"),
    ("moderate\x02intensity\x07", "moderateintensity"),
])
def test_cleaning_extractor_artifacts(raw, expected):
    assert normalize_text(raw) == expected


def test_pdf_line_end_hyphen_and_backend_is_not_agpl(tmp_path):
    path = write_pdf(tmp_path / "hyphen.pdf", ["Do 150 minutes of moderate-\nintensity activity every week."])
    [page] = parse_document(path)
    assert normalize_text(page["text"]) == "Do 150 minutes of moderate-intensity activity every week."
    import parsing.parser as parser_module
    assert "pymupdf" not in vars(parser_module) and "fitz" not in vars(parser_module)


def test_scanned_pdf_warns_and_yields_no_chunks(tmp_path, caplog):
    path = write_pdf(tmp_path / "scan.pdf", ["", ""])
    with caplog.at_level(logging.WARNING, logger="parsing.parser"):
        assert chunk_document(path, corpus_root=tmp_path) == []
    assert "no extractable text" in caplog.text and "scan.pdf" in caplog.text


def test_header_only_pdf_page_is_dropped_and_ids_do_not_shift(tmp_path):
    header = "14\nWHO guidelines on physical activity and sedentary behaviour"
    assert len(normalize_text(header)) <= MIN_PDF_CHUNK_CHARS
    path = write_pdf(tmp_path / "who_2020_physical_activity.pdf", [header, LONG])
    chunks = chunk_document(path, corpus_root=tmp_path)
    assert [(c["page"], c["id"]) for c in chunks] == [(2, "who:400:2:0")]


def test_docx_short_text_is_kept(tmp_path):
    document = Document()
    document.add_paragraph("Plank: 3 x 30 s")
    document.save(tmp_path / "short.docx")
    [chunk] = chunk_document(tmp_path / "short.docx", corpus_root=tmp_path)
    assert chunk["text"] == "Plank: 3 x 30 s" and chunk["page"] is None


def test_docx_merged_and_nested_cells(tmp_path):
    document = Document()
    table = document.add_table(rows=2, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 2))
    merged.text = "Leg day"
    table.cell(1, 0).text = "squat"
    table.cell(1, 1).text = "lunge"
    table.cell(1, 2).add_table(rows=1, cols=1).cell(0, 0).text = "nested calf raise"
    document.save(tmp_path / "plan.docx")
    [parsed] = parse_document(tmp_path / "plan.docx")
    assert parsed["text"].count("Leg day") == 1
    assert "squat\tlunge" in parsed["text"] and "nested calf raise" in parsed["text"]


@pytest.mark.parametrize("name", ["empty.pdf", "empty.docx", "broken.pdf"])
def test_unreadable_file_raises(tmp_path, name):
    path = tmp_path / name
    path.write_bytes(b"" if name.startswith("empty") else b"%PDF-1.4 not really")
    with pytest.raises(Exception):
        parse_document(path)


def test_load_chunks_skips_broken_and_lock_files(tmp_path, monkeypatch, caplog):
    from tulpar_ai.rag import index

    write_pdf(tmp_path / "good.pdf", [LONG])
    (tmp_path / "empty.docx").write_bytes(b"")
    (tmp_path / "~$good.docx").write_bytes(b"lock")
    monkeypatch.setattr(index, "CORPUS", tmp_path)
    with caplog.at_level(logging.WARNING, logger=index.__name__):
        chunks = index.load_chunks(pdf_chunk=400)
    assert [c.source for c in chunks] == ["good.pdf"]
    assert chunks[0].title == "good.pdf" and chunks[0].page == 1
    assert "empty.docx" in caplog.text and "~$good.docx" not in caplog.text


def test_index_collection_changes_with_parsing_version(tmp_path):
    from tulpar_ai.rag.embed import LocalHashEmbedder
    from tulpar_ai.rag.index import Index

    idx = Index(embedder=LocalHashEmbedder(), path=tmp_path / "qdrant")
    try:
        assert idx.collection.endswith(f"_400_p{PARSING_VERSION}")
    finally:
        idx.close()


def test_citation_without_page_has_no_page_label():
    from tulpar_ai.graph.chat import _sources

    line = _sources([{"title": "notes.docx", "page": None, "text": "Plank: 3 x 30 s"}])
    assert line.startswith("[1] notes.docx\n") and "стр." not in line
